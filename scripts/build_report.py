"""Generate the ten-page report (PDF, DOCX, Markdown) from measured outputs and screenshots."""
import json
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches
from PIL import Image as PILImage
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

from heartops.train import candidates

WIDTH = 480  # usable PDF width in points


def grid_text():
    parts = []
    for name, (_, grid) in candidates().items():
        values = '; '.join(f"{k.removeprefix('classifier__')}={v}" for k, v in grid.items())
        size = 1
        for v in grid.values():
            size *= len(v)
        parts.append(f'{name}: {values} ({size} configurations)')
    return ' | '.join(parts)


def pages(metrics, identity):
    rows = {row['name']: row for row in metrics['candidates']}
    held = metrics['heldout_metrics']
    shots = Path('screenshots')
    return [
        ('1. Project overview and architecture', [
            f"Student: {identity['student']}. Repository: {identity['repository']}.",
            'HeartOps predicts presence of heart disease from 13 routine clinical attributes in '
            'the UCI Cleveland dataset and serves the model as a monitored API. The project '
            'covers the full MLOps lifecycle: reproducible data acquisition, EDA, tuned '
            'scikit-learn pipelines, MLflow tracking and registry, automated tests, a four-stage '
            'GitHub Actions pipeline, a Docker image, a Kubernetes deployment exposed through '
            'Ingress, and Prometheus/Grafana monitoring.',
            'The diagram shows the four layers. Training produces one serialized pipeline '
            '(preprocessing + classifier) that is the only artifact the API loads, so the '
            'transformations used in training and serving are identical by construction.'],
         ['reports/figures/architecture.png']),
        ('2. Data acquisition and cleaning', [
            'Source: Janosi, Steinbrunn, Pfisterer and Detrano (1989), Heart Disease, UCI Machine '
            'Learning Repository, DOI 10.24432/C52P4X, CC BY 4.0. scripts/download_data.py (via '
            'heartops.data.acquire) downloads the official archive, extracts '
            'processed.cleveland.data, validates 303 rows and diagnosis codes 0-4, stores the '
            f"raw SHA-256 ({metrics['data_sha256'][:16]}...) and writes data/processed/heart.csv.",
            'Cleaning: "?" markers become NaN (ca: 4, thal: 2 - only 6 cells). Diagnosis 1-4 is '
            'mapped to 1 (disease present). Original UCI codes are kept (cp 1-4, slope 1-3, thal '
            '3/6/7). Missing values are deliberately not imputed on the whole dataset; imputation '
            'is learned inside each CV fold to avoid leakage.',
            f"Class balance: {metrics['class_counts']} - roughly 54/46, so accuracy is meaningful, "
            'but ROC-AUC is used for model selection.'],
         ['reports/figures/class_balance.png', 'reports/figures/missing_values.png']),
        ('3. Exploratory data analysis', [
            'Numeric densities by diagnosis show the strongest separation for maximum heart rate '
            '(mean 139 with disease vs 158 without) and ST depression (oldpeak 1.6 vs 0.6). '
            'Cholesterol and resting blood pressure overlap heavily.',
            'Per-category disease rates are striking: asymptomatic chest pain (cp=4) has a 73% '
            'disease rate versus 18-30% for other types; reversible thallium defect (thal=7) 76% '
            'versus 22% for normal (thal=3). These motivate one-hot encoding rather than treating '
            'codes as ordinal numbers.'],
         ['reports/figures/numeric_by_target.png', 'reports/figures/categorical_rates.png']),
        ('4. Feature engineering', [
            'Numeric features (age, trestbps, chol, thalach, oldpeak, ca) receive median '
            'imputation then StandardScaler. Categorical codes (sex, cp, fbs, restecg, exang, '
            'slope, thal) receive most-frequent imputation then one-hot encoding with '
            'handle_unknown="ignore". A ColumnTransformer and the classifier form one sklearn '
            'Pipeline (src/heartops/model.py), which is fitted inside every CV fold.',
            'Strongest absolute correlations with the target: thal 0.53, ca 0.46, exang 0.43, '
            'oldpeak 0.42, thalach -0.42, cp 0.41. No features were dropped: the dataset is small, '
            'and regularisation/tree depth handle weak features. EDA is descriptive only and was '
            'not used on held-out data.'],
         ['reports/figures/correlations.png']),
        ('5. Model development and tuning', [
            f"Split: stratified 80/20 with seed {metrics['seed']} - {metrics['train_rows']} "
            f"training and {metrics['test_rows']} held-out patients. 5-fold stratified CV on "
            'training patients only, scoring accuracy, precision, recall and ROC-AUC; refit on '
            'ROC-AUC.',
            f'Grids - {grid_text()}.',
            *[f"{n}: CV ROC-AUC {r['cv_roc_auc']:.4f} (std {r['cv_roc_auc_std']:.4f}), accuracy "
              f"{r['cv_accuracy']:.4f}, precision {r['cv_precision']:.4f}, recall "
              f"{r['cv_recall']:.4f}; best {r['params']}." for n, r in rows.items()],
            f"Selected: {metrics['selected']} (highest mean CV ROC-AUC). A strongly regularised "
            'linear model wins on 242 patients; the forest overfits more (higher train-fold AUC) '
            'without better validation AUC. The held-out set never influenced the choice.'],
         ['reports/figures/model_comparison.png']),
        ('6. Evaluation and experiment tracking', [
            'Held-out (61 patients, threshold 0.5): ' + ', '.join(
                f'{k} {v:.4f}' for k, v in held.items()) + '. Recall is prioritised clinically '
            '(few missed positives); with 61 patients these estimates carry wide uncertainty.',
            'MLflow (SQLite backend, experiment "heartops-cleveland"): one parent run per model '
            'family with best parameters, CV metrics, grid CSV, out-of-fold ROC and confusion '
            'plots, EDA figures and the fitted pipeline; one nested child run per grid '
            'configuration (14 trials); and a final run with held-out metrics, registered as '
            '"heartops-classifier" v1 in the Model Registry.'],
         ['reports/figures/confusion_matrix.png', 'reports/figures/roc_curve.png',
          shots / '02b_mlflow_compare_charts.png']),
        ('7. Packaging, reproducibility and CI/CD', [
            'Packaging: artifacts/model.joblib holds the complete pipeline; the same model is '
            'logged in MLflow format. requirements.txt pins every direct dependency; seeds (42) '
            'and data SHA-256 are recorded in metrics.json and MLflow tags. Clean setup: python '
            '-m venv .venv; pip install -r requirements.txt; pip install --no-deps -e .; python '
            '-m heartops.train.',
            'Tests (pytest, 16): dataset shape/missingness, provenance, EDA output, scaling and '
            'encoding, grid validity, serialization equivalence, API validation and metrics, '
            'inference CLI and a ROC-AUC >= 0.80 quality gate.',
            'GitHub Actions (.github/workflows/ci.yml) runs four dependent jobs: lint (ruff) -> '
            'train + test (JUnit XML, step summary, report) -> Docker build + HTTP smoke test '
            '-> kind deploy + ingress verification. Any failure stops later jobs; logs and '
            'artifacts (model, MLflow store, test XML, container and k8s evidence) are uploaded. '
            'Run #1 on GitHub passed all four jobs in 7m 8s.'],
         [shots / '01_lint_and_unit_tests.png', shots / '10b_github_actions_summary.png']),
        ('8. Model serving and containerization', [
            'FastAPI (src/heartops/api.py) loads the pipeline at startup. POST /predict accepts a '
            'JSON patient (Pydantic-validated ranges and codes; unknown fields rejected; ca/thal '
            'nullable) and returns prediction, disease_probability, confidence (probability of '
            'the predicted class) and threshold. /health, /ready, /metrics and /docs support '
            'operations. Invalid input returns HTTP 422 without echoing patient data.',
            'The Dockerfile uses python:3.12-slim and runs as non-root UID 10001. The image was '
            'built and tested locally: /ready returned model_loaded=true and /predict returned a '
            'prediction for examples/patient.json.'],
         [shots / '04_docker_build_and_predict.png']),
        ('9. Production deployment and monitoring', [
            'deploy/api.yaml: Deployment (2 replicas, resource limits, non-root securityContext, '
            'readiness/liveness probes, Prometheus annotations) and a LoadBalancer Service; '
            'deploy/ingress.yaml exposes it via ingress-nginx. scripts/deploy_kind.sh creates a '
            'kind cluster, loads the image, deploys and verifies /predict through the ingress at '
            'http://localhost:8081 (on kind the LoadBalancer IP stays pending; it is assigned on '
            'Docker Desktop, minikube tunnel or a cloud cluster).',
            'Monitoring: every request is logged as JSON (request id, route, status, latency; no '
            'patient data). Prometheus scrapes /metrics; the provisioned Grafana dashboard shows '
            'request rate, 4xx/5xx rate, p50/p95 latency, predictions per class and mean '
            'predicted probability.'],
         [shots / '05_kubernetes_deployment.png', shots / '09_grafana_dashboard.png']),
        ('10. Submission, access and limitations', [
            f"Repository: {identity['repository']}. Local access: docker compose up -d --build "
            '(API :8000, Prometheus :9090, Grafana :3000) or ./scripts/deploy_kind.sh (ingress '
            ':8081). Evidence is in screenshots/, artifacts/ and reports/verification.md.',
            'Notebooks: notebooks/01_eda, 02_training, 03_inference (executed). Inference CLI: '
            'scripts/predict.py (local model or --url for the deployed API).',
            'Limitations: 303 historical patients from one centre, no external validation, no '
            'probability calibration, no drift or outcome monitoring, and no API authentication. '
            'Future work: external validation, calibration, confidence intervals, drift metrics, '
            'image registry with signed digests and authenticated ingress. This is a coursework '
            'model, not a clinical tool.'],
         [shots / '08_prometheus_targets.png']),
    ]


def pdf_images(paths, max_height):
    """Images scaled to the page width; two narrow plots are placed side by side."""
    paths = [Path(p) for p in paths]
    sizes = [PILImage.open(p).size for p in paths]
    pair = len(paths) >= 2 and all(w / h < 1.8 for w, h in sizes[:2])
    flow = []
    if pair:
        cells = []
        for p, (w, h) in zip(paths[:2], sizes[:2]):
            scale = min((WIDTH / 2 - 6) / w, max_height / h)
            cells.append(Image(str(p), width=w * scale, height=h * scale))
        table = Table([cells], colWidths=[WIDTH / 2] * 2)
        table.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'MIDDLE')]))
        flow.append(table)
        paths, sizes = paths[2:], sizes[2:]
    for p, (w, h) in zip(paths, sizes):
        scale = min(WIDTH / w, max_height / h)
        flow.extend([Spacer(1, 6), Image(str(p), width=w * scale, height=h * scale)])
    return flow


def main():
    report = Path('reports')
    metrics = json.loads(Path('artifacts/metrics.json').read_text())
    identity_file = report / 'submission.json'
    if not identity_file.exists():
        identity_file.write_text(json.dumps({'student': 'ADD YOUR NAME AND ID',
                                            'repository': 'ADD YOUR REPOSITORY URL'}, indent=2))
    identity = json.loads(identity_file.read_text())
    styles = getSampleStyleSheet()
    doc = Document()
    pdf, markdown = [], []
    for index, (title, paragraphs, figures) in enumerate(pages(metrics, identity)):
        figures = [f for f in figures if Path(f).exists()]
        if index:
            doc.add_page_break()
            pdf.append(PageBreak())
        doc.add_heading(title, 1)
        pdf.append(Paragraph(escape(title), styles['Heading1']))
        markdown.append('## ' + title)
        for paragraph in paragraphs:
            doc.add_paragraph(paragraph)
            pdf.extend([Paragraph(escape(paragraph), styles['BodyText']), Spacer(1, 6)])
            markdown.append(paragraph)
        text_height = sum(len(p) for p in paragraphs) / 95 * 13 + 60
        budget = max(160, 700 - text_height)
        per_image = budget / max(1, len(figures) - (1 if len(figures) >= 2 else 0))
        pdf.extend(pdf_images(figures, per_image - 12))
        for figure in figures:
            doc.add_picture(str(figure), width=Inches(6))
            markdown.append(f'![{Path(figure).stem}](../{Path(figure).as_posix()})')
    doc.save(report / 'assignment_report.docx')
    SimpleDocTemplate(str(report / 'assignment_report.pdf'), pagesize=A4, topMargin=36,
                      bottomMargin=36, leftMargin=57, rightMargin=57).build(pdf)
    (report / 'assignment_report.md').write_text('# HeartOps - MLOps Assignment 1 report\n\n'
                                                 + '\n\n'.join(markdown) + '\n')


if __name__ == '__main__':
    main()
